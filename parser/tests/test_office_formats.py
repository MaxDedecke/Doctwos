"""Excel, PowerPoint, OpenDocument, HTML und das Alt-Format .doc: Textextraktion für Ordner, WebDAV und Upload."""

import zipfile

import pytest

from connectors import folder
from connectors.confluence import _extract_attachment_text, _is_supported_mime
from connectors.office import UnsupportedFormat, extract_legacy_doc, html_to_text


def test_xlsx_sheets_become_headed_blocks(tmp_path):
    import openpyxl

    path = tmp_path / "preise.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Preise"
    sheet.append(["Artikel", "Preis"])
    sheet.append(["Schraube", 0.25])
    workbook.create_sheet("Leer")
    workbook.save(path)

    text = folder._extract_text(str(path))
    assert "## Blatt: Preise" in text and "Schraube | 0.25" in text and "## Blatt: Leer" in text


def test_pptx_text_tables_and_notes_are_read(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "folien.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Architektur"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Dienst"
    table.cell(0, 1).text = "Port"
    slide.notes_slide.notes_text_frame.text = "Nur intern zeigen"
    presentation.save(path)

    text = folder._extract_text(str(path))
    assert "## Folie 1" in text and "Architektur" in text and "Dienst | Port" in text and "Notizen: Nur intern zeigen" in text


def test_odt_paragraphs_are_read(tmp_path):
    path = tmp_path / "brief.odt"
    content = (
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>'
        "<text:h>Titel</text:h><text:p>Erster <text:span>Absatz</text:span></text:p></office:text></office:body></office:document-content>"
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("content.xml", content)
    assert folder._extract_text(str(path)) == "Titel\nErster Absatz"


def test_html_is_stripped_to_text_with_table_cells_kept():
    html = "<html><head><title>x</title><style>p{}</style></head><body><h1>Betrieb</h1><p>Start &amp; Stopp</p><script>alert(1)</script><table><tr><td>a</td><td>b</td></tr></table></body></html>"
    text = html_to_text(html)
    assert "Betrieb" in text and "Start & Stopp" in text and "alert" not in text
    assert "| a | b" in text


def test_csv_is_read_with_windows_encoding(tmp_path):
    path = tmp_path / "liste.csv"
    path.write_bytes("Name;Ort\nMüller;Köln\n".encode("cp1252"))
    assert "Müller;Köln" in folder._extract_text(str(path))


def test_legacy_doc_is_not_guessed_but_renamed_docx_works(tmp_path):
    binary = tmp_path / "alt.doc"
    binary.write_bytes(b"\xd0\xcf\x11\xe0 binary word 97")
    with pytest.raises(UnsupportedFormat, match="als .docx speichern"):
        folder._extract_text(str(binary))

    import docx

    renamed = tmp_path / "neu.doc"
    document = docx.Document()
    document.add_paragraph("Inhalt")
    document.save(renamed)
    assert "Inhalt" in extract_legacy_doc(str(renamed), folder.extract_docx_text)


def test_confluence_attachments_cover_office_formats(tmp_path):
    import openpyxl

    assert _is_supported_mime("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    path = tmp_path / "a.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["Wert", 7])
    workbook.save(path)
    text = _extract_attachment_text(path.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert text and "Wert | 7" in text
    assert _extract_attachment_text(b"\xd0\xcf\x11\xe0 word97", "application/msword") is None
    assert "Straße" in _extract_attachment_text("Straße".encode("cp1252"), "text/plain")
